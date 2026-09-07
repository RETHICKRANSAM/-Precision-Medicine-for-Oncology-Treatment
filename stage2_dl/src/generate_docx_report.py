import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(BASE_DIR, "reports")
DOCX_PATH = os.path.join(REPORTS_DIR, "STAGE2_DATA_ENGINEER_REPORT.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 114, 178) # Teal / Blue
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(50, 50, 50)
    return p

def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(20, 20, 20)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_callout(doc, title, text, bg_hex="EBF3F9", border_hex="0072B2"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r_t = p.add_run(title)
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0, 80, 140)
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r_txt = p2.add_run(text)
    r_txt.font.name = 'Calibri'
    r_txt.font.size = Pt(10.5)
    r_txt.font.italic = True
    r_txt.font.color.rgb = RGBColor(40, 40, 40)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def generate_report():
    doc = docx.Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # -------------------------------------------------------------
    # 1. TITLE PAGE
    # -------------------------------------------------------------
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(36)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("PERSONALIZED PRECISION ONCOLOGY\nDATA ENGINEER REPORT")
    r.font.name = 'Arial'
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor(27, 54, 93)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(12)
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Role 1 of Stage 02 — Deep Learning Data Pipeline Engineering")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0, 114, 178)

    add_callout(
        doc,
        "ACADEMIC SYNTHETIC DATASET DISCLAIMER",
        "This report and its associated dataset (1,000 raw records, 750 images) represent a SYNTHETIC / MOCK oncology dataset created strictly for academic project development, demonstration, and viva presentation. No real patient health information (PHI) or clinical diagnostic images were used. Synthetic data must not be claimed as real clinical data."
    )

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(24)
    p_meta.paragraph_format.line_spacing = 1.3
    
    runs_meta = [
        ("Project Title: ", "Personalized Precision Medicine for Oncology Treatment Optimization\n"),
        ("Stage Objective: ", "Stage 02 Deep Learning Data Pipeline & Multi-Modal Dataset Engineering\n"),
        ("Prediction Objectives: ", "1. Histopathology Image Classification (CNN)\n2. Temporal ctDNA & Biomarker Progression Prediction (LSTM / Transformer)\n"),
        ("Role Responsibility: ", "Data Engineer (Role 1)\n"),
        ("Intended Audience: ", "Academic Evaluation Committee, Technical Viva Panel, Project Supervisors, & Deep Learning Team\n"),
        ("Document Version: ", "Stage 02 Final Release (v2.0)\n"),
        ("Date: ", "September 2026")
    ]
    for lbl, val in runs_meta:
        r1 = p_meta.add_run(lbl)
        r1.font.name = 'Arial'
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(27, 54, 93)
        r2 = p_meta.add_run(val)
        r2.font.name = 'Calibri'
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(50, 50, 50)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. TABLE OF CONTENTS
    # -------------------------------------------------------------
    add_heading_1(doc, "2. Table of Contents")
    toc_items = [
        ("1. Title Page & Academic Disclaimer", "1"),
        ("2. Table of Contents", "2"),
        ("3. Data Engineer Role in Deep Learning", "3"),
        ("4. Core Responsibilities & Pipeline Scope", "4"),
        ("5. End-to-End Deep Learning Data Pipeline Architecture", "5"),
        ("6. Multi-Modal Dataset Handling & Specifications", "6"),
        ("7. Systematic Data Cleaning Implementation", "8"),
        ("8. Data Outputs & Handoff Directory Structure", "10"),
        ("9. Core Data Engineering Decisions & Rationale", "11"),
        ("10. Deep Learning Team Handoff Protocol", "13"),
        ("11. Patient-Level Data Leakage Prevention Strategy", "14"),
        ("12. Technical Viva Questions and Answers (20 Q&A)", "15"),
        ("13. Data Engineer's Contribution — Simple Language Explanation", "20")
    ]
    
    tbl_toc = doc.add_table(rows=len(toc_items)+1, cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_toc.autofit = False
    
    hdr_cells = tbl_toc.rows[0].cells
    hdr_cells[0].width = Inches(5.2)
    hdr_cells[1].width = Inches(1.3)
    set_cell_background(hdr_cells[0], "1B365D")
    set_cell_background(hdr_cells[1], "1B365D")
    
    r0 = hdr_cells[0].paragraphs[0].add_run("Section / Topic")
    r0.font.bold = True; r0.font.color.rgb = RGBColor(255, 255, 255)
    r1 = hdr_cells[1].paragraphs[0].add_run("Page Ref")
    r1.font.bold = True; r1.font.color.rgb = RGBColor(255, 255, 255)
    
    for idx, (sec, pg) in enumerate(toc_items, 1):
        row_cells = tbl_toc.rows[idx].cells
        row_cells[0].width = Inches(5.2)
        row_cells[1].width = Inches(1.3)
        if idx % 2 == 0:
            set_cell_background(row_cells[0], "F8F9FA")
            set_cell_background(row_cells[1], "F8F9FA")
        row_cells[0].paragraphs[0].add_run(sec)
        row_cells[1].paragraphs[0].add_run(pg)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 3. DATA ENGINEER ROLE
    # -------------------------------------------------------------
    add_heading_1(doc, "3. Data Engineer Role in Deep Learning")
    add_body_paragraph(doc, "In Stage 02 of the Personalized Precision Oncology project, the Data Engineer serves as the foundational backbone for all Deep Learning workflows. While traditional Machine Learning (Stage 01) relies on structured tabular feature matrices, Deep Learning architectures (Convolutional Neural Networks and Transformers/LSTMs) demand high-throughput, multi-modal data ingestion pipelines capable of pairing unstructured high-resolution medical imagery with longitudinal molecular biomarkers.")
    
    add_heading_2(doc, "3.1 Why Data Engineering is Critical for Deep Learning")
    add_body_paragraph(doc, "Deep Learning models are notoriously sensitive to raw data corruption, unstandardized image dimensions, label mismatch, and temporal sequence disorder. Without robust Data Engineering:", "Garbage In, Garbage Out: ")
    add_body_paragraph(doc, "CNN classifiers will suffer from catastrophic gradient collapse or label hallucination if fed broken file paths, unverified RGB stain ranges, or mismatched categorical annotations.")
    add_body_paragraph(doc, "Recurrent (LSTM) and Attention-based (Transformer) models require strict temporal alignment across multi-timepoint longitudinal visits. Unsorted or missing timepoint intervals corrupt hidden state cell propagation.")
    add_body_paragraph(doc, "Patient-level data leakage occurs when image slices or sequential encounters of the same patient bleed across training and evaluation splits, leading to artificially inflated, non-generalizable validation metrics.")

    add_heading_2(doc, "3.2 Key Pillars of Stage 02 Data Engineering")
    pillars = [
        ("Multi-Modal Integration: ", "Harmonizing unstructured histopathology slides, CT DICOM slice references, MRI sequence scans (T1, T2, FLAIR, DWI, ADC), and longitudinal ctDNA/protein marker measurements into unified relational schemas."),
        ("Physical File & Path Integrity: ", "Verifying 100% of physical image assets on disk, flagging non-existent references, and insulating training loops from file IO errors."),
        ("Categorical & Text Standardizations: ", "Stripping dirty unit suffixes ('cm3', 'copies/mL', 'slices', 'days', '%') and mapping inconsistent clinician shorthand ('M', 'female', 'stage i', 'MALIGNANT') into normalized canonical taxonomies."),
        ("Sequence Ordering & Structuring: ", "Constructing temporal matrices sorted by Patient_ID and Biomarker_Timepoint_Days to empower sequence modeling teams.")
    ]
    for p_title, p_desc in pillars:
        add_body_paragraph(doc, p_desc, bold_prefix=p_title)

    # -------------------------------------------------------------
    # 4. RESPONSIBILITIES
    # -------------------------------------------------------------
    add_heading_1(doc, "4. Core Responsibilities & Pipeline Scope")
    add_body_paragraph(doc, "The Data Engineer (Role 1) executed a comprehensive 9-phase operational workflow to transform dirty, raw multi-modal oncology records into pristine, production-ready inputs for downstream computer vision and sequence modeling teams.")
    
    resp_table = [
        ("Phase 1: Ingestion & Profiling", "Loaded 1,000 raw encounter records containing 35 heterogenous clinical columns; generated 01_raw_data_profile.txt."),
        ("Phase 2: Image Asset Generation", "Synthesized 350 Microscopy Histopathology images, 200 CT slice images, and 200 MRI sequence images across 13 distinct sub-class folders."),
        ("Phase 3: Physical Path Validation", "Scanned disk paths for all image IDs, identifying 11 broken/missing image references in raw metadata."),
        ("Phase 4: Categorical Normalization", "Standardized Cancer Stage, Sex, Histopathology Label, Progression Status, and Progression Risk taxonomies."),
        ("Phase 5: Contamination Cleaning", "Purged embedded text units from numeric fields via regex and imputed invalid/outlier values (Age, Tumor Volume)."),
        ("Phase 6: Date Format Unification", "Converted mixed date strings (MM/DD/YYYY, DD-MM-YYYY) into standardized ISO 8601 YYYY-MM-DD formats."),
        ("Phase 7: Duplicate Purging", "Identified and removed 20 exact duplicate encounter rows, maintaining 980 pristine unique rows."),
        ("Phase 8: Specialized Handoff Export", "Produced dedicated image_metadata.csv for CNN team and temporal_biomarker_sequences.csv for LSTM/Transformer team."),
        ("Phase 9: Quality Reporting & Audit", "Generated 02_validation_report.txt, 03_cleaning_report.txt, and complete Word report documentation.")
    ]
    
    tbl_r = doc.add_table(rows=len(resp_table)+1, cols=2)
    tbl_r.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_r.autofit = False
    
    r_hdr = tbl_r.rows[0].cells
    r_hdr[0].width = Inches(2.2); r_hdr[1].width = Inches(4.3)
    set_cell_background(r_hdr[0], "1B365D"); set_cell_background(r_hdr[1], "1B365D")
    r_hdr[0].paragraphs[0].add_run("Phase").font.bold = True; r_hdr[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255,255,255)
    r_hdr[1].paragraphs[0].add_run("Engineering Scope & Deliverables").font.bold = True; r_hdr[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255,255,255)

    for idx, (ph, desc) in enumerate(resp_table, 1):
        row_c = tbl_r.rows[idx].cells
        row_c[0].width = Inches(2.2); row_c[1].width = Inches(4.3)
        if idx % 2 == 0:
            set_cell_background(row_c[0], "F8F9FA"); set_cell_background(row_c[1], "F8F9FA")
        row_c[0].paragraphs[0].add_run(ph)
        row_c[1].paragraphs[0].add_run(desc)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 5. DATA PIPELINE ARCHITECTURE
    # -------------------------------------------------------------
    add_heading_1(doc, "5. End-to-End Deep Learning Data Pipeline Architecture")
    add_body_paragraph(doc, "The Stage 02 Data Pipeline enforces strict immutability of raw assets while orchestrating a decoupled, deterministic flow from raw ingestion to model handoff.")
    
    pipeline_flow = (
        "RAW MULTI-MODAL DATASET (1,000 Records + 750 Images)\n"
        "  │\n"
        "  ├──► 1. VALIDATION ENGINE (src/validate_data.py)\n"
        "  │      ├── Schema & Column Integrity Checks (35 Columns)\n"
        "  │      ├── Physical Disk Image Verification (Histopathology, CT, MRI)\n"
        "  │      ├── Duplicate Row & Missing Value Identification\n"
        "  │      └── Output: reports/02_validation_report.txt\n"
        "  │\n"
        "  ├──► 2. CLEANING & TRANSFORMATION PIPELINE (src/clean_data.py)\n"
        "  │      ├── Duplicate Purging (1000 ──► 980 Clean Rows)\n"
        "  │      ├── Regex Unit Stripping ('cm3', 'copies/mL', 'slices', 'days', '%')\n"
        "  │      ├── Outlier Imputation (Age -5/145 ──► Median Age)\n"
        "  │      ├── Categorical Standardization & ISO Date Conversion\n"
        "  │      ├── Broken Image Path Removal & ID Synchronization\n"
        "  │      └── Output: reports/03_cleaning_report.txt\n"
        "  │\n"
        "  └──► 3. DECOUPLED DEEP LEARNING HANDOFF EXPORTS\n"
        "         ├──► data/processed/dl_cleaned.csv (Full Tabular Master)\n"
        "         ├──► data/processed/image_metadata.csv ──► CNN / Computer Vision Team\n"
        "         └──► data/processed/temporal_biomarker_sequences.csv ──► LSTM / Transformer Team"
    )
    add_callout(doc, "DATA PIPELINE ARCHITECTURAL FLOW", pipeline_flow, bg_hex="F4F6F9", border_hex="1B365D")

    # -------------------------------------------------------------
    # 6. DATASET HANDLING & SPECIFICATIONS
    # -------------------------------------------------------------
    add_heading_1(doc, "6. Multi-Modal Dataset Handling & Specifications")
    add_body_paragraph(doc, "The Stage 02 raw dataset contains 1,000 encounter records across 196 unique patients, combined with 750 synthetic medical images spanning 3 modalities.")
    
    add_heading_2(doc, "6.1 Image Asset Distribution")
    add_body_paragraph(doc, "1. Histopathology Microscopy Images (350 Files): 70 images per tissue class across 5 categories:", bold_prefix="Asset Breakdown: ")
    add_body_paragraph(doc, "   - Benign Tissue (70 files): Regular circular cellular structures, uniform purple nuclei.")
    add_body_paragraph(doc, "   - Malignant Tissue (70 files): Pleomorphic, crowded irregular cells, dark hyperchromatic nuclei.")
    add_body_paragraph(doc, "   - Atypical Cells (70 files): Mixed cell sizes, nuclear enlargement, mild architectural distortion.")
    add_body_paragraph(doc, "   - Necrotic Tissue (70 files): Disrupted, pale eosinophilic background with cellular debris.")
    add_body_paragraph(doc, "   - Inflammatory Tissue (70 files): Dense lymphocyte clusters, small dark immune infiltrates.")
    
    add_body_paragraph(doc, "2. CT Radiologic Images (200 Files): 256x256 grayscale cross-sectional representations:", bold_prefix="CT Scans: ")
    add_body_paragraph(doc, "   - Normal (65 files): Intact organ parenchyma without hyperdense masses.")
    add_body_paragraph(doc, "   - Tumor (70 files): Well-demarcated hyperdense neoplastic nodule.")
    add_body_paragraph(doc, "   - Progression (65 files): Large invasive tumor mass with surrounding hypodense edema.")

    add_body_paragraph(doc, "3. MRI Sequence Images (200 Files): 40 files per sequence representation:", bold_prefix="MRI Scans: ")
    add_body_paragraph(doc, "   - T1-Weighted (40 files): Dark CSF, mid-intensity gray parenchyma.")
    add_body_paragraph(doc, "   - T2-Weighted (40 files): Bright CSF, hyperintense fluid contrast.")
    add_body_paragraph(doc, "   - FLAIR (40 files): Fluid-attenuated dark CSF with hyperintense lesion signal.")
    add_body_paragraph(doc, "   - DWI (40 files): High signal intensity indicating restricted water diffusion.")
    add_body_paragraph(doc, "   - ADC Map (40 files): Corresponding hypointense ADC lesion signal.")

    add_heading_2(doc, "6.2 Temporal Biomarker Specifications")
    add_body_paragraph(doc, "Longitudinal molecular measurements are recorded per patient across 5 clinical timepoints (Day 0, Day 14, Day 28, Day 56, Day 84). Key temporal features include ctDNA_Level (copies/mL), Protein_Marker_Level (ng/mL), Tumor_Volume_cm3, and Tumor_Growth_Rate (%/day).")

    # -------------------------------------------------------------
    # 7. DATA CLEANING IMPLEMENTATION
    # -------------------------------------------------------------
    add_heading_1(doc, "7. Systematic Data Cleaning Implementation")
    add_body_paragraph(doc, "The cleaning engine (src/clean_data.py) executes a deterministic 8-stage transformation pipeline:")
    
    clean_steps = [
        ("1. Duplicate Purging", "Identified 20 duplicate encounter rows where all 35 columns matched. Purged duplicates to yield 980 clean rows."),
        ("2. Categorical Standardization", "Mapped inconsistent string entries (e.g. 'M', 'male', '  Male ' ──► 'Male'; 'stage i', 'I' ──► 'Stage I'; 'MALIGNANT', 'benign' ──► Canonical Title Case)."),
        ("3. Regex Unit Contamination Stripping", "Extracted pure float values from strings containing embedded text units ('25.4 cm3' ──► 25.4; '12.8 copies/mL' ──► 12.8; '120 slices' ──► 120; '28 days' ──► 28)."),
        ("4. Outlier & Range Correction", "Replaced impossible age values (-5, 145) and negative tumor volumes (-15.2 cm3) with dataset median values."),
        ("5. Date Format Unification", "Parsed irregular date formats (MM/DD/YYYY, DD-MM-YYYY, YYYY/MM/DD) into standardized ISO 8601 YYYY-MM-DD."),
        ("6. Image Path Disk Audit", "Cross-referenced all image path strings against physical disk storage. Flagged 11 broken paths (e.g. HIMG99999.png) and cleared invalid path references."),
        ("7. Missing Value Imputation", "Imputed missing ctDNA and Protein Marker levels with median values; imputed missing categorical labels with 'Unknown' or mode."),
        ("8. Longitudinal Sequence Sorting", "Sorted master datasets by Patient_ID and Biomarker_Timepoint_Days to enforce causal temporal ordering for sequence modeling.")
    ]
    for s_name, s_desc in clean_steps:
        add_body_paragraph(doc, s_desc, bold_prefix=f"{s_name}: ")

    # -------------------------------------------------------------
    # 8. DATA OUTPUTS
    # -------------------------------------------------------------
    add_heading_1(doc, "8. Data Outputs & Handoff Directory Structure")
    add_body_paragraph(doc, "All generated data products adhere strictly to the standardized stage2_dl directory layout:")

    dir_struct = (
        "stage2_dl/\n"
        "├── data/\n"
        "│   ├── raw/\n"
        "│   │   ├── dl_raw_1000.csv                    [1,000 raw encounter records, 35 columns]\n"
        "│   │   ├── histopathology_images/             [350 synthetic PNG images in 5 subfolders]\n"
        "│   │   ├── ct_images/                         [200 synthetic PNG images in 3 subfolders]\n"
        "│   │   └── mri_images/                        [200 synthetic PNG images in 5 subfolders]\n"
        "│   └── processed/\n"
        "│       ├── dl_cleaned.csv                     [980 cleaned master encounter records]\n"
        "│       ├── image_metadata.csv                 [346 verified image records for CNN team]\n"
        "│       └── temporal_biomarker_sequences.csv   [980 temporal sequence records for LSTM team]\n"
        "├── src/\n"
        "│   ├── create_dataset.py                      [Synthetic raw generator script]\n"
        "│   ├── validate_data.py                      [Data quality & disk audit engine]\n"
        "│   └── clean_data.py                         [Data cleaning & export pipeline]\n"
        "├── reports/\n"
        "│   ├── 01_raw_data_profile.txt                [Raw data schema & synthetic profile]\n"
        "│   ├── 02_validation_report.txt               [Validation audit findings]\n"
        "│   ├── 03_cleaning_report.txt                 [Cleaning transformations audit]\n"
        "│   └── STAGE2_DATA_ENGINEER_REPORT.docx       [Comprehensive Word documentation report]\n"
        "├── notebooks/\n"
        "│   └── stage2_data_exploration.ipynb          [Exploratory data analysis notebook]\n"
        "└── README.md                                  [Stage 02 documentation & viva guide]"
    )
    add_callout(doc, "COMPLETE STAGE 02 DIRECTORY STRUCTURE", dir_struct, bg_hex="F8F9FA", border_hex="1B365D")

    # -------------------------------------------------------------
    # 9. DATA ENGINEERING DECISIONS
    # -------------------------------------------------------------
    add_heading_1(doc, "9. Core Data Engineering Decisions & Rationale")
    
    decisions = [
        ("1. Immutability of Raw Data: ", "The raw CSV (dl_raw_1000.csv) and raw image files are treated as immutable read-only assets. Preserving raw data guarantees full auditability, reproducibility, and lineage tracing if cleaning heuristics are revised."),
        ("2. Separation of Processed Data: ", "Processed datasets are stored strictly in data/processed/. Decoupling raw ingestion from clean outputs prevents accidental overwrites and ensures clean downstream imports."),
        ("3. Decoupling Image Metadata from Temporal Sequences: ", "Exporting separate image_metadata.csv and temporal_biomarker_sequences.csv allows the Computer Vision (CNN) and Sequence (LSTM/Transformer) teams to train independently without handling irrelevant features or missing modal paths."),
        ("4. Strict Physical Disk Validation: ", "Image path strings in metadata must be validated against actual physical files on disk before handoff. Removing broken references prevents runtime FileNotFoundError crashes during PyTorch/TensorFlow DataLoader execution."),
        ("5. Mandatory Temporal Sequence Sorting: ", "Sequence models rely on causal hidden state transitions. Sorting by Patient_ID and Biomarker_Timepoint_Days guarantees that sequence windows (e.g. Day 0 -> Day 14 -> Day 28) represent true chronological progression.")
    ]
    for d_title, d_desc in decisions:
        add_body_paragraph(doc, d_desc, bold_prefix=d_title)

    # -------------------------------------------------------------
    # 10. DEEP LEARNING TEAM HANDOFF
    # -------------------------------------------------------------
    add_heading_1(doc, "10. Deep Learning Team Handoff Protocol")
    add_body_paragraph(doc, "To facilitate seamless collaboration across specialized modeling roles, the Data Engineer provides dedicated handoff specifications:")
    
    add_heading_2(doc, "10.1 CNN / Computer Vision Team Handoff")
    add_body_paragraph(doc, "data/processed/image_metadata.csv & data/raw/histopathology_images/", bold_prefix="Primary Deliverables: ")
    add_body_paragraph(doc, "Image_ID, Patient_ID, Image_Path, Label, Cancer_Type, Cancer_Stage, Tissue_Type, Tumor_Grade, Image_Quality.", bold_prefix="Key Input Columns: ")
    add_body_paragraph(doc, "Histopathology_Label (5 classes: Benign, Malignant, Atypical, Necrotic, Inflammatory).", bold_prefix="Primary Target Column: ")
    add_body_paragraph(doc, "Load image tensors using PIL/OpenCV from Image_Path. Apply PyTorch/TF transforms: Resize (224x224), Normalization (ImageNet mean/std), Random Horizontal/Vertical Flip, Color Jitter.", bold_prefix="Recommended Processing: ")

    add_heading_2(doc, "10.2 LSTM / Transformer Sequence Modeling Team Handoff")
    add_body_paragraph(doc, "data/processed/temporal_biomarker_sequences.csv", bold_prefix="Primary Deliverables: ")
    add_body_paragraph(doc, "Patient_ID, Timepoint_Days, ctDNA_Level, Protein_Marker_Level, Tumor_Volume_cm3, Tumor_Growth_Rate.", bold_prefix="Key Input Features: ")
    add_body_paragraph(doc, "Progression_Risk (3 classes: Low, Moderate, High) or Progression_Status (Progressed, Stable, Not Progressed).", bold_prefix="Primary Target Column: ")
    add_body_paragraph(doc, "Group records by Patient_ID into sequential 3D tensors of shape (Batch_Size, Sequence_Length, Feature_Dim). Apply MinMaxScaler or StandardScaler across temporal feature dimensions.", bold_prefix="Recommended Processing: ")

    # -------------------------------------------------------------
    # 11. DATA LEAKAGE PREVENTION
    # -------------------------------------------------------------
    add_heading_1(doc, "11. Patient-Level Data Leakage Prevention Strategy")
    add_body_paragraph(doc, "Data leakage occurs when information from outside the training dataset is inadvertently used to train the model, resulting in overly optimistic evaluation metrics that fail in real clinical deployment.")
    
    add_callout(
        doc,
        "CRITICAL LEAKAGE WARNING: PATIENT-LEVEL SPLITTING REQUIRED",
        "Because individual patients have multiple longitudinal encounter records (Day 0, Day 14, Day 28, Day 56, Day 84) and multiple image scans, naive random row-level train/test splits WILL contaminate the test set. Multiple timepoints of the same patient would appear in both training and test sets, causing severe data leakage!"
    )
    
    add_body_paragraph(doc, "1. Grouped Stratified K-Fold Split: Always split data using GroupKFold or GroupShuffleSplit grouped strictly by Patient_ID. This ensures that 100% of a patient's images, timepoints, and clinical notes reside exclusively in either the Train, Validation, or Test split.", bold_prefix="Prevention Protocol: ")
    add_body_paragraph(doc, "2. Scaler Imputation Isolation: Compute feature scaling parameters (mean, std, min, max) strictly on the Training set split. Apply the fitted scalers to Validation and Test sets without recalculating parameters.", bold_prefix="Scaler Protocol: ")

    # -------------------------------------------------------------
    # 12. VIVA QUESTIONS AND ANSWERS
    # -------------------------------------------------------------
    add_heading_1(doc, "12. Technical Viva Questions and Answers")
    add_body_paragraph(doc, "This section provides 20 rigorously structured Viva Q&A pairs covering key Stage 02 Data Engineering concepts for final-year project evaluations.")

    viva_qas = [
        ("Q1: Why is Data Engineering different for Deep Learning compared to standard ML?",
         "Standard ML relies primarily on flat tabular matrices. Deep Learning handles multi-modal unstructured assets (images, video, raw audio) and longitudinal sequences, requiring high-throughput data pipelines, physical image path validation, stain normalization, and temporal sequence tensor generation."),
        ("Q2: Why is image metadata separation required?",
         "Decoupling image metadata into a dedicated CSV (image_metadata.csv) allows Computer Vision teams to stream image paths, labels, and quality scores directly into PyTorch DataLoaders without dragging along heavy longitudinal tabular records or missing sequence features."),
        ("Q3: What is image data augmentation and why is it useful?",
         "Image augmentation applies artificial geometric and color transformations (rotations, flips, zooming, color jitter) to training images. It artificially increases dataset diversity, prevents CNN overfitting, and teaches the network rotation/scale invariance."),
        ("Q4: What is a Convolutional Neural Network (CNN)?",
         "A CNN is a deep neural network architecture designed for spatial grid data like images. It utilizes trainable convolutional filters to automatically extract hierarchical spatial features—from low-level edges and textures to high-level cellular structures."),
        ("Q5: What is histopathology image classification in oncology?",
         "Histopathology classification analyzes microscopic tissue slide images (typically H&E stained) to classify tissue regions into diagnostic categories such as Benign, Malignant, Atypical, Necrotic, or Inflammatory to guide cancer diagnosis."),
        ("Q6: What is temporal biomarker data in precision medicine?",
         "Temporal biomarker data consists of longitudinal molecular measurements (e.g. ctDNA levels, serum protein markers) collected over multiple clinical timepoints (e.g., Days 0, 14, 28, 56) to track treatment response and disease progression."),
        ("Q7: Why use an LSTM network for temporal biomarker modeling?",
         "Long Short-Term Memory (LSTM) networks are recurrent neural architectures equipped with memory cell gates designed to learn long-term sequential dependencies and avoid vanishing gradient problems when modeling longitudinal patient trajectories."),
        ("Q8: Why use a Transformer model instead of an LSTM for sequence data?",
         "Transformers utilize self-attention mechanisms to process sequence timepoints in parallel rather than sequentially. They capture global temporal relationships regardless of sequence length and avoid sequential bottlenecking."),
        ("Q9: What is ctDNA and why is it significant in oncology monitoring?",
         "Circulating tumor DNA (ctDNA) refers to fragmented DNA released by cancer cells into the bloodstream. Measuring ctDNA levels provides a non-invasive 'liquid biopsy' to quantify minimal residual disease and early treatment response."),
        ("Q10: Why are multiple longitudinal timepoints required per patient?",
         "A single static biomarker measurement cannot distinguish between responding, stable, or progressing tumors. Multiple timepoints reveal velocity (growth rate) and trajectory trends essential for dynamic outcome prediction."),
        ("Q11: What is an image-label mismatch and how does it impact DL models?",
         "Image-label mismatch occurs when an image file path is paired with an incorrect diagnostic label (e.g., a malignant slide labeled as benign). It introduces severe label noise, confusing gradient updates and degrading model accuracy."),
        ("Q12: How do you detect corrupted or broken image files in data engineering?",
         "By executing disk audit scripts (like src/validate_data.py) that attempt to verify file existence via os.path.exists() and load image headers using PIL.Image.open() to trap OSError/UnidentifiedImageError before training."),
        ("Q13: Why must raw data files remain untouched?",
         "Maintaining immutable raw data ensures reproducible data lineage, complete auditability, and allows data engineers to re-run or modify cleaning pipelines without data loss or permanent corruption."),
        ("Q14: What is data leakage and how do you prevent it in Stage 02?",
         "Data leakage occurs when test set information contaminates training. In Stage 02, it is prevented by performing patient-level splits (GroupKFold by Patient_ID) so all encounters/images of a patient stay in a single split."),
        ("Q15: Why must dataset splits be performed at the patient level rather than the row level?",
         "Because multiple timepoints and images belong to the same patient. Naive row-level splitting would put Day 0 of Patient A in Train and Day 14 of Patient A in Test, causing the model to memorize patient-specific features."),
        ("Q16: What is class imbalance in histopathology datasets?",
         "Class imbalance occurs when certain diagnostic classes (e.g. malignant) vastly outnumber others (e.g. necrotic). It causes models to bias towards majority classes, remediable via focal loss or oversampling."),
        ("Q17: How did you handle contaminated text in numeric CSV fields?",
         "Using regular expressions in pandas (src/clean_data.py) to strip text suffixes (e.g., '25.4 cm3' -> 25.4 float; '12.8 copies/mL' -> 12.8 float) before casting to float numeric types."),
        ("Q18: What is the primary role of a Data Engineer in a Deep Learning project?",
         "To build robust, automated pipelines that ingest, validate, clean, format, and decouple multi-modal data into high-quality, audit-ready inputs optimized for model training."),
        ("Q19: What specific datasets were handed off to the DL engineers?",
         "image_metadata.csv (plus histopathology/CT/MRI image folders) for the CNN team, and temporal_biomarker_sequences.csv for the LSTM/Transformer sequence team."),
        ("Q20: What disclaimers apply to this dataset?",
         "The dataset is 100% synthetic/mock, generated strictly for academic demonstration, pipeline development, and viva presentation. It contains no real patient data.")
    ]

    for q, a in viva_qas:
        add_body_paragraph(doc, q, bold_prefix="")
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_before = Pt(0)
        p_ans.paragraph_format.space_after = Pt(6)
        p_ans.paragraph_format.left_indent = Inches(0.2)
        r_a = p_ans.add_run(a)
        r_a.font.name = 'Calibri'
        r_a.font.size = Pt(10.5)
        r_a.font.color.rgb = RGBColor(60, 60, 60)

    # -------------------------------------------------------------
    # 13. SIMPLE LANGUAGE EXPLANATION
    # -------------------------------------------------------------
    add_heading_1(doc, "13. Data Engineer's Contribution — Simple Language Explanation")
    
    simple_summary = (
        "I prepared and organized the synthetic oncology image and biomarker data required for the Deep Learning stage. "
        "I validated image paths and labels, handled metadata quality issues, prepared temporal biomarker sequences, "
        "kept raw data untouched, and produced structured datasets that the CNN and LSTM/Transformer teams can directly use."
    )
    add_callout(doc, "DATA ENGINEER'S SUMMARY STATEMENT", simple_summary, bg_hex="EBF3F9", border_hex="0072B2")

    doc.save(DOCX_PATH)
    print(f"Generated Stage 02 Data Engineer Report at: {DOCX_PATH}")

if __name__ == "__main__":
    generate_report()
